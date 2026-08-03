"""
Servicio de exportación a Excel para movimientos y reportes.

Genera archivos Excel profesionales con formato, colores y estilos.
Usa openpyxl para máxima compatibilidad con Microsoft Excel.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO


def crear_estilo_header():
    """Retorna un diccionario con estilos para el encabezado."""
    return {
        'font': Font(name='Calibri', size=11, bold=True, color='FFFFFF'),
        'fill': PatternFill(start_color='0066CC', end_color='0066CC', fill_type='solid'),
        'alignment': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'border': Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
    }


def crear_estilo_datos():
    """Retorna un diccionario con estilos para datos normales."""
    return {
        'font': Font(name='Calibri', size=10),
        'alignment': Alignment(horizontal='left', vertical='center', wrap_text=False),
        'border': Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )
    }


def aplicar_estilo(cell, estilos):
    """Aplica múltiples estilos a una celda."""
    if 'font' in estilos:
        cell.font = estilos['font']
    if 'fill' in estilos:
        cell.fill = estilos['fill']
    if 'alignment' in estilos:
        cell.alignment = estilos['alignment']
    if 'border' in estilos:
        cell.border = estilos['border']


def ajustar_ancho_columnas(worksheet, max_width=50):
    """Ajusta automáticamente el ancho de las columnas basado en el contenido."""
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            try:
                if cell.value:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            except:
                pass

        adjusted_width = min(max_length + 2, max_width)
        worksheet.column_dimensions[column_letter].width = adjusted_width


def exportar_movimientos_excel(movimientos, filtros=None):
    """
    Exporta una lista de movimientos a Excel.

    Args:
        movimientos: Lista de objetos Movimiento (con joins)
        filtros: Diccionario con filtros aplicados (opcional)

    Returns:
        BytesIO: Archivo Excel en memoria
    """
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    # Estilos
    estilo_header = crear_estilo_header()
    estilo_datos = crear_estilo_datos()

    # ========== TÍTULO DEL REPORTE ==========
    ws.merge_cells('A1:J1')
    titulo = ws['A1']
    titulo.value = "REPORTE DE MOVIMIENTOS - JEROSMART ACTIVOS"
    titulo.font = Font(name='Calibri', size=14, bold=True, color='0066CC')
    titulo.alignment = Alignment(horizontal='center', vertical='center')

    # Fecha de generación
    ws.merge_cells('A2:J2')
    fecha_gen = ws['A2']
    fecha_gen.value = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    fecha_gen.font = Font(name='Calibri', size=10, italic=True)
    fecha_gen.alignment = Alignment(horizontal='center')

    # Filtros aplicados
    if filtros:
        fila_filtros = 3
        texto_filtros = []
        if filtros.get('q'):
            texto_filtros.append(f"Búsqueda: {filtros['q']}")
        if filtros.get('tipo'):
            texto_filtros.append(f"Tipo: {filtros['tipo']}")

        if texto_filtros:
            ws.merge_cells(f'A{fila_filtros}:J{fila_filtros}')
            filtro_cell = ws[f'A{fila_filtros}']
            filtro_cell.value = f"Filtros aplicados: {' | '.join(texto_filtros)}"
            filtro_cell.font = Font(name='Calibri', size=9, italic=True, color='666666')
            filtro_cell.alignment = Alignment(horizontal='center')

    # ========== ENCABEZADOS ==========
    fila_header = 5
    headers = [
        'ID',
        'Tipo Movimiento',
        'Fecha',
        'Usuario',
        'Estado Aprobación',
        'Cantidad Activos',
        'Valor Total (COP)',
        'Aprobado Por',
        'Fecha Aprobación',
        'Observaciones'
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=fila_header, column=col_num)
        cell.value = header
        aplicar_estilo(cell, estilo_header)

    # ========== DATOS ==========
    fila_datos = fila_header + 1

    for movimiento in movimientos:
        # Calcular valor total del movimiento
        valor_total = sum(
            ma.valor_libros_momento for ma in movimiento.activos
        ) if hasattr(movimiento, 'activos') else 0

        # Determinar estado de aprobación con color
        estado_aprobacion = movimiento.estado_aprobacion or 'Pendiente'

        datos = [
            movimiento.id,
            movimiento.tipo_movimiento,
            movimiento.fecha.strftime('%d/%m/%Y %H:%M') if movimiento.fecha else '',
            movimiento.usuario.email if movimiento.usuario else 'N/A',
            estado_aprobacion,
            len(movimiento.activos) if hasattr(movimiento, 'activos') else 0,
            f"${valor_total:,.2f}" if valor_total > 0 else '$0',
            movimiento.aprobador.email if movimiento.aprobador else '',
            movimiento.fecha_aprobacion.strftime('%d/%m/%Y') if movimiento.fecha_aprobacion else '',
            (movimiento.observaciones_generales or '')[:100]  # Limitar a 100 caracteres
        ]

        for col_num, valor in enumerate(datos, 1):
            cell = ws.cell(row=fila_datos, column=col_num)
            cell.value = valor
            aplicar_estilo(cell, estilo_datos)

            # Colorear estado de aprobación
            if col_num == 5:  # Columna de estado
                if estado_aprobacion == 'Aprobado':
                    cell.fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='155724', bold=True)
                elif estado_aprobacion == 'Rechazado':
                    cell.fill = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='721C24', bold=True)
                elif estado_aprobacion == 'Pendiente':
                    cell.fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='856404', bold=True)

        fila_datos += 1

    # ========== PIE DE PÁGINA ==========
    fila_footer = fila_datos + 1
    ws.merge_cells(f'A{fila_footer}:J{fila_footer}')
    footer = ws[f'A{fila_footer}']
    footer.value = f"Total de movimientos: {len(movimientos)}"
    footer.font = Font(name='Calibri', size=10, bold=True)
    footer.alignment = Alignment(horizontal='right')

    # Ajustar anchos de columnas
    ajustar_ancho_columnas(ws)

    # Congelar paneles (primera fila de datos)
    ws.freeze_panes = f'A{fila_header + 1}'

    # Guardar en memoria
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output


def exportar_movimiento_detallado_excel(movimiento):
    """
    Exporta un movimiento individual con todos sus detalles a Excel.

    Args:
        movimiento: Objeto Movimiento con relaciones cargadas

    Returns:
        BytesIO: Archivo Excel en memoria
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"Movimiento {movimiento.id}"

    estilo_header = crear_estilo_header()
    estilo_datos = crear_estilo_datos()

    # ========== INFORMACIÓN GENERAL ==========
    fila = 1
    ws.merge_cells(f'A{fila}:D{fila}')
    titulo = ws[f'A{fila}']
    titulo.value = f"DETALLE DE MOVIMIENTO #{movimiento.id}"
    titulo.font = Font(name='Calibri', size=14, bold=True, color='0066CC')
    titulo.alignment = Alignment(horizontal='center')

    fila += 2
    campos_generales = [
        ('Tipo de Movimiento:', movimiento.tipo_movimiento),
        ('Fecha:', movimiento.fecha.strftime('%d/%m/%Y %H:%M') if movimiento.fecha else ''),
        ('Usuario Creador:', movimiento.usuario.email if movimiento.usuario else 'N/A'),
        ('Estado de Aprobación:', movimiento.estado_aprobacion),
        ('Aprobado Por:', movimiento.aprobador.email if movimiento.aprobador else ''),
        ('Fecha de Aprobación:', movimiento.fecha_aprobacion.strftime('%d/%m/%Y') if movimiento.fecha_aprobacion else ''),
    ]

    for campo, valor in campos_generales:
        ws[f'A{fila}'] = campo
        ws[f'A{fila}'].font = Font(bold=True)
        ws[f'B{fila}'] = valor
        fila += 1

    # ========== ACTIVOS INVOLUCRADOS ==========
    fila += 2
    ws.merge_cells(f'A{fila}:F{fila}')
    subtitulo = ws[f'A{fila}']
    subtitulo.value = "ACTIVOS INVOLUCRADOS"
    subtitulo.font = Font(name='Calibri', size=12, bold=True, color='0066CC')
    subtitulo.alignment = Alignment(horizontal='center')

    fila += 1
    headers_activos = ['Placa', 'Nombre', 'Marca', 'Modelo', 'Valor Libros', 'Ubicación Origen']

    for col_num, header in enumerate(headers_activos, 1):
        cell = ws.cell(row=fila, column=col_num)
        cell.value = header
        aplicar_estilo(cell, estilo_header)

    fila += 1
    valor_total = 0

    for ma in movimiento.activos:
        activo = ma.activo
        datos_activo = [
            activo.placa_codigo_interno,
            activo.nombre_activo,
            activo.marca or '',
            activo.modelo or '',
            f"${ma.valor_libros_momento:,.2f}",
            ma.ubicacion_origen or ''
        ]

        for col_num, valor in enumerate(datos_activo, 1):
            cell = ws.cell(row=fila, column=col_num)
            cell.value = valor
            aplicar_estilo(cell, estilo_datos)

        valor_total += ma.valor_libros_momento
        fila += 1

    # Total
    fila += 1
    ws.merge_cells(f'A{fila}:D{fila}')
    total_cell = ws[f'A{fila}']
    total_cell.value = "VALOR TOTAL:"
    total_cell.font = Font(bold=True)
    total_cell.alignment = Alignment(horizontal='right')

    valor_cell = ws[f'E{fila}']
    valor_cell.value = f"${valor_total:,.2f}"
    valor_cell.font = Font(bold=True, color='0066CC')

    # Ajustar columnas
    ajustar_ancho_columnas(ws)

    # Guardar
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output


def exportar_estadisticas_excel(stats):
    """
    Exporta estadísticas del dashboard a Excel.

    Args:
        stats: Diccionario con estadísticas del dashboard

    Returns:
        BytesIO: Archivo Excel en memoria
    """
    wb = Workbook()

    # ========== HOJA 1: RESUMEN ==========
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"

    estilo_header = crear_estilo_header()

    # Título
    ws_resumen.merge_cells('A1:C1')
    titulo = ws_resumen['A1']
    titulo.value = "ESTADÍSTICAS DE MOVIMIENTOS"
    titulo.font = Font(name='Calibri', size=14, bold=True, color='0066CC')
    titulo.alignment = Alignment(horizontal='center')

    # Estadísticas principales
    fila = 3
    estadisticas = [
        ('Total de Movimientos:', stats.get('total_movimientos', 0)),
        ('Movimientos Este Mes:', stats.get('movimientos_mes_actual', 0)),
        ('Pendientes de Aprobación:', stats.get('pendientes_aprobacion', 0)),
        ('Aprobados Este Mes:', stats.get('aprobados_mes', 0)),
        ('Valor Total Mes (COP):', f"${stats.get('valor_total_mes', 0):,.2f}"),
    ]

    for etiqueta, valor in estadisticas:
        ws_resumen[f'A{fila}'] = etiqueta
        ws_resumen[f'A{fila}'].font = Font(bold=True)
        ws_resumen[f'B{fila}'] = valor
        ws_resumen[f'B{fila}'].font = Font(size=12, color='0066CC')
        fila += 1

    # ========== HOJA 2: MOVIMIENTOS POR TIPO ==========
    ws_tipos = wb.create_sheet("Por Tipo")
    ws_tipos['A1'] = "Tipo de Movimiento"
    ws_tipos['B1'] = "Cantidad"
    aplicar_estilo(ws_tipos['A1'], estilo_header)
    aplicar_estilo(ws_tipos['B1'], estilo_header)

    for i, (tipo, cant) in enumerate(zip(stats.get('tipos_labels', []), stats.get('tipos_data', [])), 2):
        ws_tipos[f'A{i}'] = tipo
        ws_tipos[f'B{i}'] = cant

    ajustar_ancho_columnas(ws_tipos)

    # ========== HOJA 3: ACTIVOS MÁS MOVIDOS ==========
    ws_activos = wb.create_sheet("Activos Más Movidos")
    ws_activos['A1'] = "Placa"
    ws_activos['B1'] = "Nombre"
    ws_activos['C1'] = "Cantidad de Movimientos"

    for cell in [ws_activos['A1'], ws_activos['B1'], ws_activos['C1']]:
        aplicar_estilo(cell, estilo_header)

    for i, activo in enumerate(stats.get('activos_mas_movidos', []), 2):
        ws_activos[f'A{i}'] = activo['placa']
        ws_activos[f'B{i}'] = activo['nombre']
        ws_activos[f'C{i}'] = activo['cantidad']

    ajustar_ancho_columnas(ws_activos)

    # Guardar
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output


def exportar_inventario_excel(activos, filtros=None):
    """
    Exporta el inventario de activos a Excel.

    Args:
        activos: Lista de objetos Activo (con joins)
        filtros: Diccionario con filtros aplicados (opcional)

    Returns:
        BytesIO: Archivo Excel en memoria
    """
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario de Activos"

    # Estilos
    estilo_header = crear_estilo_header()
    estilo_datos = crear_estilo_datos()

    # ========== TÍTULO DEL REPORTE ==========
    ws.merge_cells('A1:M1')
    titulo = ws['A1']
    titulo.value = "INVENTARIO GENERAL DE ACTIVOS - JEROSMART"
    titulo.font = Font(name='Calibri', size=14, bold=True, color='0066CC')
    titulo.alignment = Alignment(horizontal='center', vertical='center')

    # Fecha de generación
    ws.merge_cells('A2:M2')
    fecha_gen = ws['A2']
    fecha_gen.value = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    fecha_gen.font = Font(name='Calibri', size=10, italic=True)
    fecha_gen.alignment = Alignment(horizontal='center')

    # Filtros aplicados
    if filtros:
        fila_filtros = 3
        texto_filtros = []
        if filtros.get('q'):
            texto_filtros.append(f"Búsqueda: {filtros['q']}")
        if filtros.get('tipo'):
            texto_filtros.append(f"Tipo: {filtros['tipo']}")
        if filtros.get('clase'):
            texto_filtros.append(f"Clase ID: {filtros['clase']}")
        if filtros.get('fecha_desde'):
            texto_filtros.append(f"Desde: {filtros['fecha_desde']}")
        if filtros.get('fecha_hasta'):
            texto_filtros.append(f"Hasta: {filtros['fecha_hasta']}")

        if texto_filtros:
            ws.merge_cells(f'A{fila_filtros}:M{fila_filtros}')
            filtro_cell = ws[f'A{fila_filtros}']
            filtro_cell.value = f"Filtros aplicados: {' | '.join(texto_filtros)}"
            filtro_cell.font = Font(name='Calibri', size=9, italic=True, color='666666')
            filtro_cell.alignment = Alignment(horizontal='center')

    # ========== ENCABEZADOS ==========
    fila_header = 5
    headers = [
        'Placa/Código',
        'Nombre',
        'Marca',
        'Modelo',
        'Serie',
        'Tipo Propiedad',
        'Clase',
        'Estado',
        'Ubicación',
        'Valor Comercial (COP)',
        'Fecha Ingreso',
        'Responsable',
        'Observaciones'
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=fila_header, column=col_num)
        cell.value = header
        aplicar_estilo(cell, estilo_header)

    # ========== DATOS ==========
    fila_datos = fila_header + 1
    valor_total = 0

    for activo in activos:
        # Obtener nombre de clase si existe
        nombre_clase = ''
        if hasattr(activo, 'nombre_clase') and activo.nombre_clase:
            nombre_clase = activo.nombre_clase
        elif activo.clase:
            nombre_clase = activo.clase.nombre_clase

        # Obtener nombre del funcionario responsable
        responsable = ''
        if hasattr(activo, 'funcionario_responsable') and activo.funcionario_responsable:
            responsable = activo.funcionario_responsable
        elif activo.funcionario:
            responsable = f"{activo.funcionario.nombres} {activo.funcionario.apellidos}"

        # Para activos ajenos, mostrar el propietario
        if activo.tipo_propiedad == 'Ajeno' and activo.propietario_ajeno:
            responsable = f"{activo.propietario_ajeno} (Propietario)"

        datos = [
            activo.placa_codigo_interno,
            activo.nombre_activo,
            activo.marca or '',
            activo.modelo or '',
            activo.serie or '',
            activo.tipo_propiedad or 'Propio',
            nombre_clase,
            activo.estado or 'Operativo',
            activo.ubicacion or '',
            activo.valor_comercial if activo.valor_comercial else 0,
            activo.created_at.strftime('%d/%m/%Y') if activo.created_at else '',
            responsable,
            (activo.observaciones or '')[:100]  # Limitar a 100 caracteres
        ]

        for col_num, valor in enumerate(datos, 1):
            cell = ws.cell(row=fila_datos, column=col_num)

            # Formatear valor comercial como moneda
            if col_num == 10:  # Columna de valor comercial
                if valor and valor > 0:
                    cell.value = valor
                    cell.number_format = '$#,##0.00'
                    valor_total += valor
                else:
                    cell.value = '$0.00'
            else:
                cell.value = valor

            aplicar_estilo(cell, estilo_datos)

            # Colorear tipo de propiedad
            if col_num == 6:  # Columna de tipo propiedad
                if valor == 'Propio':
                    cell.fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='155724', bold=True)
                elif valor == 'Ajeno':
                    cell.fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='856404', bold=True)

            # Colorear estado
            if col_num == 8:  # Columna de estado
                if valor == 'Operativo':
                    cell.fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='155724', bold=True)
                elif valor in ['En Mantenimiento', 'En préstamo']:
                    cell.fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='856404', bold=True)
                elif valor in ['Dañado', 'Dado de Baja', 'Retirado']:
                    cell.fill = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='721C24', bold=True)

        fila_datos += 1

    # ========== PIE DE PÁGINA ==========
    fila_footer = fila_datos + 1

    # Total de activos
    ws.merge_cells(f'A{fila_footer}:I{fila_footer}')
    footer = ws[f'A{fila_footer}']
    footer.value = f"Total de activos: {len(activos)}"
    footer.font = Font(name='Calibri', size=10, bold=True)
    footer.alignment = Alignment(horizontal='right')

    # Valor total
    ws[f'J{fila_footer}'] = valor_total
    ws[f'J{fila_footer}'].number_format = '$#,##0.00'
    ws[f'J{fila_footer}'].font = Font(name='Calibri', size=10, bold=True, color='0066CC')

    # Ajustar anchos de columnas
    ajustar_ancho_columnas(ws)

    # Congelar paneles (primera fila de datos)
    ws.freeze_panes = f'A{fila_header + 1}'

    # Guardar en memoria
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output
